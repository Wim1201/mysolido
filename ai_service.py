"""
MySolido AI Service — Lokale AI-assistent op pod-data
Gebruikt Ollama voor LLM en embeddings, ChromaDB voor vector-opslag.
Standalone bruikbaar als module of via MySolido Flask-interface.
"""

import os
import json
import logging
from datetime import datetime

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Configuration (overridable via environment variables)
# ---------------------------------------------------------------------------

OLLAMA_BASE_URL = os.getenv("OLLAMA_BASE_URL", "http://127.0.0.1:11434")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "tinyllama")
OLLAMA_EMBED_MODEL = os.getenv("OLLAMA_EMBED_MODEL", "nomic-embed-text")

# AI Provider niveaus (nu en toekomst)
# Niveau 1: "local"     — Ollama voor alles (standaard)
# Niveau 2: "hybrid"    — Lokaal indexeren + Claude API voor antwoorden
# Niveau 3: "cloud_ocr" — Niveau 2 + Claude Vision voor OCR (toekomst, NIET gebouwd)
#   - extract_text() krijgt ook een provider-keuze
#   - Bij "cloud_ocr": afbeeldingen en gescande PDF's gaan naar Claude Vision
#   - extract_ocr() wordt extract_ocr_local() + extract_ocr_claude()
#   - Instelling: drie radiobuttons in plaats van twee

AI_PROVIDER = os.getenv("AI_PROVIDER", "local")  # "local" of "hybrid"
ANTHROPIC_API_KEY = os.getenv("ANTHROPIC_API_KEY", "")
ANTHROPIC_MODEL = os.getenv("ANTHROPIC_MODEL", "claude-sonnet-4-20250514")

CHROMA_PERSIST_DIR = os.getenv(
    "CHROMA_PERSIST_DIR",
    os.path.join(os.path.dirname(os.path.abspath(__file__)), ".chroma"),
)

SUPPORTED_EXTENSIONS = {
    ".jsonld": "json",
    ".json": "json",
    ".txt": "text",
    ".md": "text",
    ".markdown": "text",
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".xls": "xlsx",
    ".csv": "text",
    ".png": "ocr",
    ".jpg": "ocr",
    ".jpeg": "ocr",
    ".gif": "ocr",
    ".tiff": "ocr",
    ".bmp": "ocr",
}

# ---------------------------------------------------------------------------
# Optional dependency checks
# ---------------------------------------------------------------------------

_MISSING_DEPS = {}

try:
    import chromadb
except ImportError:
    chromadb = None
    _MISSING_DEPS["chromadb"] = "pip install chromadb"

try:
    import requests as _requests
except ImportError:
    _requests = None
    _MISSING_DEPS["requests"] = "pip install requests"

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None
    _MISSING_DEPS["PyPDF2"] = "pip install PyPDF2"

try:
    from docx import Document as _DocxDocument
except ImportError:
    _DocxDocument = None
    _MISSING_DEPS["python-docx"] = "pip install python-docx"

try:
    import openpyxl
except ImportError:
    openpyxl = None
    _MISSING_DEPS["openpyxl"] = "pip install openpyxl"

try:
    from PIL import Image as _PILImage
    import pytesseract
except ImportError:
    _PILImage = None
    pytesseract = None
    _MISSING_DEPS["pytesseract"] = "pip install pytesseract  (+ Tesseract OCR op systeem)"

try:
    from pdf2image import convert_from_path as _convert_from_path
except ImportError:
    _convert_from_path = None
    _MISSING_DEPS["pdf2image"] = "pip install pdf2image  (+ poppler-utils op systeem)"


def get_missing_dependencies():
    """Return dict of missing optional dependencies and install hints."""
    return dict(_MISSING_DEPS)


# ---------------------------------------------------------------------------
# Text extraction pipeline
# ---------------------------------------------------------------------------

def extract_text(file_path):
    """Extract readable text from any supported file type."""
    ext = os.path.splitext(file_path)[1].lower()
    file_type = SUPPORTED_EXTENSIONS.get(ext)
    if file_type is None:
        return ""
    try:
        if file_type == "json":
            return _extract_json(file_path)
        if file_type == "text":
            return _extract_text_file(file_path)
        if file_type == "pdf":
            return _extract_pdf(file_path)
        if file_type == "docx":
            return _extract_docx(file_path)
        if file_type == "xlsx":
            return _extract_xlsx(file_path)
        if file_type == "ocr":
            return _extract_ocr(file_path)
    except Exception as exc:
        logger.warning("Text extraction failed for %s: %s", file_path, exc)
    return ""


def _json_to_readable(obj, prefix=""):
    """Flatten a JSON/JSON-LD object into readable key-value lines."""
    lines = []
    if isinstance(obj, dict):
        for key, value in obj.items():
            if key == "@context":
                continue
            full_key = f"{prefix}{key}" if not prefix else f"{prefix}.{key}"
            if isinstance(value, (dict, list)):
                lines.extend(_json_to_readable(value, full_key))
            else:
                lines.append(f"{full_key}: {value}")
    elif isinstance(obj, list):
        for i, item in enumerate(obj):
            lines.extend(_json_to_readable(item, f"{prefix}[{i}]"))
    else:
        lines.append(f"{prefix}: {obj}")
    return lines


def _extract_json(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    return "\n".join(_json_to_readable(data))


def _extract_text_file(file_path):
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _extract_pdf(file_path):
    if PyPDF2 is None:
        logger.warning("PyPDF2 not installed — skipping PDF: %s", file_path)
        return ""
    text = ""
    with open(file_path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for page in reader.pages:
            page_text = page.extract_text() or ""
            text += page_text + "\n"
    # Fall back to OCR for scanned PDFs
    if len(text.strip()) < 50:
        ocr_text = _extract_ocr_from_pdf(file_path)
        if ocr_text:
            text = ocr_text
    return text.strip()


def _extract_docx(file_path):
    if _DocxDocument is None:
        logger.warning("python-docx not installed — skipping DOCX: %s", file_path)
        return ""
    doc = _DocxDocument(file_path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_xlsx(file_path):
    if openpyxl is None:
        logger.warning("openpyxl not installed — skipping XLSX: %s", file_path)
        return ""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(c) for c in row if c is not None)
            if row_text.strip():
                parts.append(row_text)
    wb.close()
    return "\n".join(parts)


def _extract_ocr(file_path):
    if _PILImage is None or pytesseract is None:
        logger.warning("pytesseract/Pillow not installed — skipping OCR: %s", file_path)
        return ""
    try:
        image = _PILImage.open(file_path)
        return pytesseract.image_to_string(image, lang="nld+eng").strip()
    except Exception as exc:
        logger.warning("OCR failed for %s: %s", file_path, exc)
        return ""


def _extract_ocr_from_pdf(file_path):
    if _convert_from_path is None or pytesseract is None:
        return ""
    try:
        images = _convert_from_path(file_path, dpi=200)
        text = ""
        for img in images:
            text += pytesseract.image_to_string(img, lang="nld+eng") + "\n"
        return text.strip()
    except Exception as exc:
        logger.warning("PDF OCR failed for %s: %s", file_path, exc)
        return ""


# ---------------------------------------------------------------------------
# Text chunking
# ---------------------------------------------------------------------------

def chunk_text(text, max_words=500, overlap_words=50):
    """Split text into overlapping word-based chunks."""
    words = text.split()
    if len(words) <= max_words:
        return [text]
    chunks = []
    start = 0
    while start < len(words):
        end = start + max_words
        chunks.append(" ".join(words[start:end]))
        start = end - overlap_words
    return chunks


# ---------------------------------------------------------------------------
# ChromaDB helpers
# ---------------------------------------------------------------------------

def _require_chromadb():
    if chromadb is None:
        raise RuntimeError(
            "chromadb is niet geinstalleerd. Installeer met: pip install chromadb"
        )


def get_chroma_client():
    _require_chromadb()
    return chromadb.PersistentClient(path=CHROMA_PERSIST_DIR)


def get_or_create_collection(client):
    return client.get_or_create_collection(
        name="mysolido_documents",
        metadata={"hnsw:space": "cosine"},
    )


# ---------------------------------------------------------------------------
# Ollama helpers
# ---------------------------------------------------------------------------

def _require_requests():
    if _requests is None:
        raise RuntimeError("requests is niet geinstalleerd. Installeer met: pip install requests")


def get_embedding(text):
    """Get embedding vector from Ollama."""
    _require_requests()
    try:
        resp = _requests.post(
            f"{OLLAMA_BASE_URL}/api/embeddings",
            json={"model": OLLAMA_EMBED_MODEL, "prompt": text},
            timeout=30,
        )
        if resp.status_code == 200:
            return resp.json()["embedding"]
    except Exception as exc:
        logger.warning("Embedding request failed: %s", exc)
    return None


def check_ollama_status():
    """Check whether Ollama is running and which models are available."""
    _require_requests()
    try:
        resp = _requests.get(f"{OLLAMA_BASE_URL}/api/tags", timeout=5)
        if resp.status_code != 200:
            return {"running": False, "models": [], "error": "Ollama antwoordt niet"}

        models = [m["name"] for m in resp.json().get("models", [])]
        llm_base = OLLAMA_MODEL.split(":")[0]
        embed_base = OLLAMA_EMBED_MODEL.split(":")[0]

        return {
            "running": True,
            "models": models,
            "has_llm": OLLAMA_MODEL in models or any(llm_base in m for m in models),
            "has_embed": OLLAMA_EMBED_MODEL in models or any(embed_base in m for m in models),
            "llm_model": OLLAMA_MODEL,
            "embed_model": OLLAMA_EMBED_MODEL,
        }
    except _requests.exceptions.ConnectionError:
        return {"running": False, "models": [], "error": "Ollama is niet geinstalleerd of draait niet"}
    except Exception as exc:
        return {"running": False, "models": [], "error": str(exc)}


def check_ai_status():
    """Controleer status van alle AI-componenten."""
    ollama = check_ollama_status()
    provider = AI_PROVIDER
    claude_configured = bool(ANTHROPIC_API_KEY)

    return {
        "provider": provider,
        "ollama": ollama,
        "claude_configured": claude_configured,
        "claude_model": ANTHROPIC_MODEL if claude_configured else None,
        # Ollama is altijd nodig (voor embeddings), ook in hybrid-modus
        "ready": ollama.get("running", False) and ollama.get("has_embed", False),
    }


# ---------------------------------------------------------------------------
# Indexing service
# ---------------------------------------------------------------------------

def index_document(collection, file_path, pod_data_path):
    """Index a single document into ChromaDB."""
    text = extract_text(file_path)
    if not text or len(text.strip()) < 10:
        return False

    relative_path = os.path.relpath(file_path, pod_data_path).replace("\\", "/")
    chunks = chunk_text(text, max_words=500, overlap_words=50)

    for i, chunk_content in enumerate(chunks):
        doc_id = f"{relative_path}::chunk_{i}"
        embedding = get_embedding(chunk_content)
        if embedding:
            collection.upsert(
                ids=[doc_id],
                embeddings=[embedding],
                documents=[chunk_content],
                metadatas=[
                    {
                        "file_path": relative_path,
                        "chunk_index": i,
                        "file_type": os.path.splitext(file_path)[1].lower(),
                        "indexed_at": datetime.utcnow().isoformat(),
                    }
                ],
            )
    return True


def index_all_documents(pod_data_path, progress_callback=None):
    """Walk the pod directory tree and index every supported file."""
    _require_chromadb()
    client = get_chroma_client()
    collection = get_or_create_collection(client)

    indexed = 0
    skipped = 0
    errors = 0

    for root, dirs, files in os.walk(pod_data_path):
        # Skip hidden/internal directories
        dirs[:] = [d for d in dirs if not d.startswith(".")]

        for filename in files:
            file_path = os.path.join(root, filename)
            ext = os.path.splitext(filename)[1].lower()

            if ext not in SUPPORTED_EXTENSIONS:
                skipped += 1
                continue

            try:
                if index_document(collection, file_path, pod_data_path):
                    indexed += 1
                else:
                    skipped += 1
            except Exception as exc:
                errors += 1
                msg = f"Fout bij {filename}: {str(exc)[:100]}"
                logger.warning(msg)
                if progress_callback:
                    progress_callback(msg)

    return {"indexed": indexed, "skipped": skipped, "errors": errors}


def get_index_stats():
    """Return statistics about the current ChromaDB index."""
    try:
        client = get_chroma_client()
        collection = get_or_create_collection(client)
        count = collection.count()
        return {"total_chunks": count, "status": "ready" if count > 0 else "empty"}
    except Exception:
        return {"total_chunks": 0, "status": "not_initialized"}


# ---------------------------------------------------------------------------
# Query service
# ---------------------------------------------------------------------------

def query_documents(question, n_results=5):
    """Find the most relevant document chunks for a question."""
    _require_chromadb()
    client = get_chroma_client()
    collection = get_or_create_collection(client)

    if collection.count() == 0:
        return []

    question_embedding = get_embedding(question)
    if not question_embedding:
        return []

    results = collection.query(
        query_embeddings=[question_embedding],
        n_results=n_results,
        include=["documents", "metadatas", "distances"],
    )

    return [
        {"text": doc, "file_path": meta["file_path"], "distance": dist}
        for doc, meta, dist in zip(
            results["documents"][0],
            results["metadatas"][0],
            results["distances"][0],
        )
    ]


_SYSTEM_PROMPT = (
    "Je bent de MySolido AI-assistent. Je helpt de gebruiker met vragen "
    "over documenten in zijn persoonlijke datakluis.\n\n"
    "Regels:\n"
    "- Beantwoord vragen alleen op basis van de aangeleverde context\n"
    "- Als het antwoord niet in de context staat, zeg dat eerlijk\n"
    "- Verwijs naar het bronbestand bij je antwoord\n"
    "- Antwoord in het Nederlands tenzij de gebruiker Engels spreekt\n"
    "- Wees beknopt en direct\n"
    "- Geef nooit medisch, juridisch of financieel advies — verwijs naar professionals"
)


def generate_answer(question, context_docs):
    """Genereer antwoord via de geconfigureerde provider."""
    if AI_PROVIDER == "hybrid" and ANTHROPIC_API_KEY:
        return generate_answer_claude(question, context_docs)
    return generate_answer_ollama(question, context_docs)


def generate_answer_ollama(question, context_docs):
    """Genereer antwoord via Ollama (lokaal)."""
    _require_requests()

    context = "\n\n---\n\n".join(
        f"[Bron: {doc['file_path']}]\n{doc['text']}" for doc in context_docs
    )
    user_prompt = f"Context uit de datakluis:\n\n{context}\n\nVraag: {question}"

    try:
        resp = _requests.post(
            f"{OLLAMA_BASE_URL}/api/chat",
            json={
                "model": OLLAMA_MODEL,
                "messages": [
                    {"role": "system", "content": _SYSTEM_PROMPT},
                    {"role": "user", "content": user_prompt},
                ],
                "stream": False,
            },
            timeout=300,
        )
        if resp.status_code == 200:
            return resp.json()["message"]["content"]
        return "Er ging iets mis bij het genereren van een antwoord."
    except _requests.exceptions.ConnectionError:
        return "Ollama is niet bereikbaar. Zorg dat Ollama draait (ollama serve)."
    except _requests.exceptions.Timeout:
        return "Het genereren van een antwoord duurde te lang. Probeer een kortere vraag."
    except Exception as exc:
        logger.error("generate_answer_ollama failed: %s", exc)
        return "Er ging iets mis bij het genereren van een antwoord."


def generate_answer_claude(question, context_docs):
    """Genereer antwoord via Claude API (hybride modus)."""
    _require_requests()

    context = "\n\n---\n\n".join(
        f"[Bron: {doc['file_path']}]\n{doc['text']}" for doc in context_docs
    )

    try:
        resp = _requests.post(
            "https://api.anthropic.com/v1/messages",
            headers={
                "x-api-key": ANTHROPIC_API_KEY,
                "anthropic-version": "2023-06-01",
                "content-type": "application/json",
            },
            json={
                "model": ANTHROPIC_MODEL,
                "max_tokens": 1024,
                "system": _SYSTEM_PROMPT,
                "messages": [
                    {"role": "user", "content": f"Context uit de datakluis:\n\n{context}\n\nVraag: {question}"}
                ],
            },
            timeout=60,
        )

        if resp.status_code == 200:
            return resp.json()["content"][0]["text"]
        if resp.status_code == 401:
            return "Ongeldige API-key. Controleer je Anthropic API-key in de instellingen."
        if resp.status_code == 429:
            return "Te veel verzoeken. Probeer het over een minuut opnieuw."
        return f"Fout bij Claude API (status {resp.status_code}). Probeer het opnieuw of schakel over naar lokaal."
    except _requests.exceptions.ConnectionError:
        return "Kan geen verbinding maken met de Claude API. Controleer je internetverbinding."
    except _requests.exceptions.Timeout:
        return "De Claude API reageert niet. Probeer het opnieuw."
    except Exception as exc:
        logger.error("generate_answer_claude failed: %s", exc)
        return f"Onverwachte fout: {str(exc)[:200]}"


def ask(question):
    """Main entry point: ask a question and get an answer with sources."""
    docs = query_documents(question, n_results=5)

    if not docs:
        return {
            "answer": (
                "Ik heb geen relevante documenten gevonden in je kluis. "
                "Heb je de index al opgebouwd?"
            ),
            "sources": [],
        }

    answer = generate_answer(question, docs)
    sources = list(dict.fromkeys(doc["file_path"] for doc in docs))

    return {"answer": answer, "sources": sources}
